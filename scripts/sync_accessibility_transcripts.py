from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
document_path = ROOT / "Interview_Transcripts_Accessibility.docx"
doc = Document(document_path)
sections = [
    ("Interview 01", "Interview 02", ROOT / "interview-transcript.txt"),
    ("Interview 02", "Interview 03", ROOT / "interview-02-transcript.txt"),
]
def remove_paragraph(p):
    p._element.getparent().remove(p._element)

for start_label, end_label, source_path in sections:
    paras = list(doc.paragraphs)
    start_index = next(i for i,p in enumerate(paras) if p.text.strip() == start_label)
    end_index = next(i for i,p in enumerate(paras) if i > start_index and p.text.strip() == end_label)
    ending = paras[end_index]
    for p in paras[start_index+1:end_index]:
        remove_paragraph(p)
    for line in source_path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        para = ending.insert_paragraph_before(line.strip())
        para.style = doc.styles["Normal"]
        if line.startswith(("INTERVIEWER:", "PARTICIPANT", "Interviewee:", "Interviewer:", "01.", "02.", "03.", "04.", "05.", "06.", "07.", "08.", "09.", "10.", "Follow-up")):
            pass
# Adjust introduction without changing interviews 03 and 04.
for para in doc.paragraphs[:5]:
    if para.text.startswith("Research ethics note:"):
        para.text = ("Research ethics note: Participants are identified by number rather than name. "
            "Interviews 01 and 02 are caption-assisted, interviewer-edited working transcripts; "
            "uncertain phrases are marked and they have not been independently verified word for word. "
            "Interviews 03 and 04 contain edited or paraphrased responses from automatic captions; "
            "consult the original recordings before quoting exact wording.")
        break

doc.save(document_path)
print("Updated interviews 01–02; preserved interviews 03–04.")
